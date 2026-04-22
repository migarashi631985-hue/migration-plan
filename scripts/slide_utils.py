"""Common helpers for building training slide decks (pptx)."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

FACILITY_NAME = "みんなの介護ゆうあい"
JP_FONT = "Yu Gothic"

PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0xC0, 0x50, 0x4D)
LIGHT = RGBColor(0xE7, 0xEE, 0xF7)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x55, 0x55)


@dataclass
class DeckMeta:
    month: str
    theme: str
    subtitle: str
    method: str
    required: bool = False
    filename: str = ""
    objectives: list[str] = field(default_factory=list)


def _set_font(run, *, size: int, bold: bool = False, color: RGBColor = DARK) -> None:
    run.font.name = JP_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Also set East Asian font via XML
    rPr = run._r.get_or_add_rPr()
    from pptx.oxml.ns import qn
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {"typeface": JP_FONT})
        rPr.append(ea)
    else:
        ea.set("typeface", JP_FONT)


def _add_textbox(slide, left, top, width, height, text, *, size=18,
                 bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Cm(0.1)
    tf.margin_right = Cm(0.1)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        _set_font(run, size=size, bold=bold, color=color)
    return box


def _add_rect(slide, left, top, width, height, *, fill=PRIMARY, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def _page_header(slide, title: str, month: str):
    _add_rect(slide, Cm(0), Cm(0), Cm(33.87), Cm(1.6), fill=PRIMARY)
    _add_textbox(slide, Cm(0.6), Cm(0.25), Cm(24), Cm(1.1),
                 title, size=22, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    _add_textbox(slide, Cm(27), Cm(0.3), Cm(6.5), Cm(1),
                 f"{month}　{FACILITY_NAME}", size=11,
                 color=RGBColor(0xFF, 0xFF, 0xFF), align=PP_ALIGN.RIGHT)


def _page_footer(slide, page_no: int, total: int):
    _add_rect(slide, Cm(0), Cm(18.4), Cm(33.87), Cm(0.6), fill=LIGHT)
    _add_textbox(slide, Cm(0.6), Cm(18.45), Cm(20), Cm(0.5),
                 FACILITY_NAME, size=9, color=GREY)
    _add_textbox(slide, Cm(27), Cm(18.45), Cm(6.3), Cm(0.5),
                 f"{page_no} / {total}", size=9, color=GREY, align=PP_ALIGN.RIGHT)


def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Cm(33.87)   # 16:9 widescreen
    prs.slide_height = Cm(19.05)
    return prs


def add_title_slide(prs: Presentation, meta: DeckMeta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, Cm(0), Cm(0), Cm(33.87), Cm(19.05), fill=PRIMARY)
    _add_rect(slide, Cm(0), Cm(6.5), Cm(33.87), Cm(5.5),
              fill=RGBColor(0xFF, 0xFF, 0xFF))
    _add_textbox(slide, Cm(1.5), Cm(1.6), Cm(30), Cm(1.2),
                 f"{meta.month}　職員研修",
                 size=22, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    required_tag = "【法定必須】" if meta.required else ""
    _add_textbox(slide, Cm(1.5), Cm(6.9), Cm(30.8), Cm(2.0),
                 f"{required_tag}{meta.theme}",
                 size=40, bold=True, color=PRIMARY)
    _add_textbox(slide, Cm(1.5), Cm(9.8), Cm(30.8), Cm(1.5),
                 meta.subtitle, size=20, color=DARK)
    _add_textbox(slide, Cm(1.5), Cm(13.5), Cm(30.8), Cm(1.0),
                 FACILITY_NAME, size=22, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    _add_textbox(slide, Cm(1.5), Cm(15.0), Cm(30.8), Cm(0.8),
                 f"実施方法：{meta.method}", size=14,
                 color=RGBColor(0xEE, 0xEE, 0xEE))
    _add_textbox(slide, Cm(1.5), Cm(16.0), Cm(30.8), Cm(0.8),
                 "講師：＿＿＿＿＿　　日時：＿＿＿＿＿", size=12,
                 color=RGBColor(0xDD, 0xDD, 0xDD))


def add_objective_slide(prs: Presentation, meta: DeckMeta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _page_header(slide, "本日の目的と到達目標", meta.month)
    _add_textbox(slide, Cm(1.0), Cm(2.2), Cm(32), Cm(1.0),
                 "この研修が終わったとき、受講者は次ができるようになります。",
                 size=16, color=GREY)
    y = 4.0
    for i, obj in enumerate(meta.objectives, start=1):
        _add_rect(slide, Cm(1.0), Cm(y), Cm(1.2), Cm(1.2), fill=ACCENT)
        _add_textbox(slide, Cm(1.0), Cm(y), Cm(1.2), Cm(1.2),
                     str(i), size=20, bold=True,
                     color=RGBColor(0xFF, 0xFF, 0xFF),
                     align=PP_ALIGN.CENTER)
        _add_textbox(slide, Cm(2.6), Cm(y + 0.05), Cm(30), Cm(1.2),
                     obj, size=18, color=DARK)
        y += 1.6


def add_section_slide(prs: Presentation, meta: DeckMeta, title: str,
                       bullets: Iterable[str], *, note: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _page_header(slide, title, meta.month)
    y = 2.4
    for bullet in bullets:
        _add_rect(slide, Cm(1.2), Cm(y + 0.45), Cm(0.35), Cm(0.35),
                  fill=PRIMARY)
        _add_textbox(slide, Cm(1.9), Cm(y), Cm(31), Cm(1.4),
                     bullet, size=18, color=DARK)
        y += 1.4
    if note:
        _add_rect(slide, Cm(1.0), Cm(16.2), Cm(31.9), Cm(1.7),
                  fill=LIGHT)
        _add_textbox(slide, Cm(1.4), Cm(16.35), Cm(31), Cm(1.5),
                     f"▶ {note}", size=13, color=PRIMARY, bold=True)


def add_twocol_slide(prs: Presentation, meta: DeckMeta, title: str,
                     left_title: str, left_items: Iterable[str],
                     right_title: str, right_items: Iterable[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _page_header(slide, title, meta.month)
    # left
    _add_rect(slide, Cm(1.0), Cm(2.4), Cm(15.5), Cm(1.1), fill=PRIMARY)
    _add_textbox(slide, Cm(1.2), Cm(2.55), Cm(15), Cm(0.9),
                 left_title, size=16, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    _add_rect(slide, Cm(1.0), Cm(3.5), Cm(15.5), Cm(12.5), fill=LIGHT)
    y = 3.8
    for item in left_items:
        _add_textbox(slide, Cm(1.3), Cm(y), Cm(14.9), Cm(1.2),
                     f"・{item}", size=14, color=DARK)
        y += 1.2
    # right
    _add_rect(slide, Cm(17.3), Cm(2.4), Cm(15.5), Cm(1.1), fill=ACCENT)
    _add_textbox(slide, Cm(17.5), Cm(2.55), Cm(15), Cm(0.9),
                 right_title, size=16, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    _add_rect(slide, Cm(17.3), Cm(3.5), Cm(15.5), Cm(12.5),
              fill=RGBColor(0xFD, 0xF0, 0xEF))
    y = 3.8
    for item in right_items:
        _add_textbox(slide, Cm(17.6), Cm(y), Cm(14.9), Cm(1.2),
                     f"・{item}", size=14, color=DARK)
        y += 1.2


def add_case_slide(prs: Presentation, meta: DeckMeta, title: str,
                   situation: str, questions: Iterable[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _page_header(slide, title, meta.month)
    _add_rect(slide, Cm(1.0), Cm(2.4), Cm(31.9), Cm(1.0), fill=ACCENT)
    _add_textbox(slide, Cm(1.3), Cm(2.5), Cm(31), Cm(0.9),
                 "事例（グループ討議用）", size=15, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF))
    _add_rect(slide, Cm(1.0), Cm(3.5), Cm(31.9), Cm(6.0),
              fill=RGBColor(0xFD, 0xF0, 0xEF))
    _add_textbox(slide, Cm(1.4), Cm(3.7), Cm(31.1), Cm(5.8),
                 situation, size=16, color=DARK)
    _add_textbox(slide, Cm(1.0), Cm(10.0), Cm(31.9), Cm(0.8),
                 "考えてみましょう", size=15, bold=True, color=PRIMARY)
    y = 11.0
    for i, q in enumerate(questions, start=1):
        _add_textbox(slide, Cm(1.3), Cm(y), Cm(31.1), Cm(1.1),
                     f"Q{i}.　{q}", size=14, color=DARK)
        y += 1.2


def add_summary_slide(prs: Presentation, meta: DeckMeta,
                      points: Iterable[str], next_action: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _page_header(slide, "まとめ・今日から実践すること", meta.month)
    _add_textbox(slide, Cm(1.0), Cm(2.2), Cm(32), Cm(1.0),
                 "本日の要点", size=18, bold=True, color=PRIMARY)
    y = 3.4
    for p in points:
        _add_rect(slide, Cm(1.1), Cm(y + 0.4), Cm(0.5), Cm(0.5),
                  fill=ACCENT)
        _add_textbox(slide, Cm(2.0), Cm(y), Cm(30.5), Cm(1.3),
                     p, size=16, color=DARK)
        y += 1.3
    _add_rect(slide, Cm(1.0), Cm(14.0), Cm(31.9), Cm(3.2), fill=LIGHT)
    _add_textbox(slide, Cm(1.4), Cm(14.2), Cm(31.1), Cm(0.8),
                 "明日からの行動目標", size=15, bold=True, color=PRIMARY)
    _add_textbox(slide, Cm(1.4), Cm(15.1), Cm(31.1), Cm(2.0),
                 next_action, size=14, color=DARK)


def add_check_slide(prs: Presentation, meta: DeckMeta,
                    items: Iterable[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _page_header(slide, "自己チェック（受講後に記入）", meta.month)
    _add_textbox(slide, Cm(1.0), Cm(2.2), Cm(32), Cm(1.0),
                 "できている=◎　ときどき=○　できていない=△",
                 size=14, color=GREY)
    y = 3.5
    for i, item in enumerate(items, start=1):
        _add_rect(slide, Cm(29.0), Cm(y + 0.1), Cm(3.7), Cm(1.0),
                  fill=LIGHT)
        _add_textbox(slide, Cm(29.0), Cm(y + 0.25), Cm(3.7), Cm(0.8),
                     "◎ ・ ○ ・ △", size=14, color=PRIMARY,
                     align=PP_ALIGN.CENTER)
        _add_textbox(slide, Cm(1.0), Cm(y + 0.2), Cm(27.5), Cm(1.0),
                     f"{i}.　{item}", size=14, color=DARK)
        y += 1.25


def add_qa_slide(prs: Presentation, meta: DeckMeta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, Cm(0), Cm(0), Cm(33.87), Cm(19.05), fill=PRIMARY)
    _add_textbox(slide, Cm(1.5), Cm(6.0), Cm(30.8), Cm(3),
                 "質疑応答", size=60, bold=True,
                 color=RGBColor(0xFF, 0xFF, 0xFF),
                 align=PP_ALIGN.CENTER)
    _add_textbox(slide, Cm(1.5), Cm(10.0), Cm(30.8), Cm(1.5),
                 "ご清聴ありがとうございました", size=22,
                 color=RGBColor(0xEE, 0xEE, 0xEE),
                 align=PP_ALIGN.CENTER)
    _add_textbox(slide, Cm(1.5), Cm(16.0), Cm(30.8), Cm(0.8),
                 f"{meta.month}　{meta.theme}　／　{FACILITY_NAME}",
                 size=12,
                 color=RGBColor(0xDD, 0xDD, 0xDD),
                 align=PP_ALIGN.CENTER)


def finalise(prs: Presentation):
    # Apply footer with page numbers now that total count is known
    total = len(prs.slides)
    for idx, slide in enumerate(prs.slides, start=1):
        if idx == 1 or idx == total:
            continue  # title and qa slides
        _page_footer(slide, idx, total)


def save(prs: Presentation, path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
