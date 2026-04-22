"""Write narration scripts to a Word document (.docx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


JP_FONT = "Yu Gothic"


def _set_jp_font(run, *, size: int = 11, bold: bool = False,
                 color: RGBColor | None = None) -> None:
    run.font.name = JP_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), JP_FONT)
    rFonts.set(qn("w:ascii"), JP_FONT)
    rFonts.set(qn("w:hAnsi"), JP_FONT)


def _add_para(doc, text: str, *, size=11, bold=False,
              color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_jp_font(run, size=size, bold=bold, color=color)
    return p


def build_narration_doc(
    out_path: Path,
    *,
    facility: str,
    month: str,
    theme: str,
    subtitle: str,
    slides: list[tuple[str, str]],
) -> None:
    """Generate a Word document with narration script per slide.

    Each entry in ``slides`` is ``(slide_title, narration_text)``.
    """
    doc = Document()

    # Section setup: A4, narrower margins for readability
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
    GREY = RGBColor(0x55, 0x55, 0x55)

    _add_para(doc, facility, size=11, color=GREY)
    _add_para(doc, f"{month}　{theme}", size=22, bold=True, color=PRIMARY,
              space_after=4)
    _add_para(doc, subtitle, size=12, color=GREY, space_after=6)
    _add_para(doc, "ナレーション原稿（発表者用台本）", size=12, bold=True,
              space_after=12)

    for i, (title, text) in enumerate(slides, start=1):
        _add_para(doc, f"スライド {i:02d}　{title}",
                  size=13, bold=True, color=PRIMARY, space_after=3)
        # Narration body – split sentences onto separate lines for reading ease
        for sentence in _split_sentences(text):
            _add_para(doc, sentence, size=11, space_after=2)
        _add_para(doc, "", size=11, space_after=6)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def _split_sentences(text: str) -> list[str]:
    """Break narration on Japanese sentence terminators for legibility."""
    out: list[str] = []
    buf = ""
    for ch in text:
        buf += ch
        if ch in "。！？":
            out.append(buf.strip())
            buf = ""
    if buf.strip():
        out.append(buf.strip())
    return out
